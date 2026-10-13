import csv
import io
import json
import logging
import os
import time

from fastapi import APIRouter, Body, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import JSONResponse, RedirectResponse, StreamingResponse
from pydantic import BaseModel
from sqlalchemy import text as sql_text

from app.database.session import get_db_session

logger = logging.getLogger(__name__)
from collections import defaultdict as _dd
from datetime import datetime, timezone
from io import StringIO
from typing import Any, Dict, List, Optional, Set
from uuid import UUID

from fastapi import Request, status
from jsonschema import Draft202012Validator
from jsonschema import exceptions as schema_exceptions
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from starlette.concurrency import run_in_threadpool

from app.models.db_connection import DBConnectionConfig
from app.routers.dependencies import get_current_user
from app.routers.realtime import realtime_broadcast
from app.services.db_connection_service import (
    DBConnectionService,
    SessionExpiredError,
    SessionNotFoundError,
)

# Imports para generación de scripts de migración
from app.services.migration_script_service import (
    HashNotFoundError,
    MigrationScriptService,
    MigrationScriptServiceError,
)
from app.utils.db_adapters import (
    DBConnectionError,
    MySQLAdapter,
    OracleAdapter,
    PostgreSQLAdapter,
    SQLServerAdapter,
)
from app.utils.db_adapters.type_mappers import TypeMapper
from app.utils.error_codes import ErrorCode, make_error
from app.utils.i18n import DEFAULT_LOCALE
from app.utils.redis_client import get_redis_client
from config import make_translator

router = APIRouter(prefix="/api/mappings", tags=["mappings"])

log = logging.getLogger("import")

# ---------------------------------------------------------------------------
# Type-compatibility matrix for column mapping validation.
# Keys and values are NORMALISED type names (uppercase, no parameters).
# A mapping is allowed if dest_norm is in _TYPE_COMPATIBILITY[origin_norm].
# ---------------------------------------------------------------------------
_TYPE_COMPATIBILITY: Dict[str, List[str]] = {
    # Integer family
    "INTEGER":          ["INTEGER", "BIGINT", "SMALLINT", "REAL", "FLOAT",
                         "DOUBLE PRECISION", "NUMERIC", "DECIMAL", "TEXT", "VARCHAR", "CHAR"],
    "BIGINT":           ["BIGINT", "INTEGER", "SMALLINT", "REAL", "FLOAT",
                         "DOUBLE PRECISION", "NUMERIC", "DECIMAL", "TEXT", "VARCHAR", "CHAR"],
    "SMALLINT":         ["SMALLINT", "INTEGER", "BIGINT", "REAL", "FLOAT",
                         "DOUBLE PRECISION", "NUMERIC", "DECIMAL", "TEXT", "VARCHAR", "CHAR"],
    # Float / exact-numeric family
    "REAL":             ["REAL", "FLOAT", "DOUBLE PRECISION", "NUMERIC", "DECIMAL",
                         "INTEGER", "BIGINT", "SMALLINT", "TEXT", "VARCHAR", "CHAR"],
    "FLOAT":            ["FLOAT", "REAL", "DOUBLE PRECISION", "NUMERIC", "DECIMAL",
                         "INTEGER", "BIGINT", "SMALLINT", "TEXT", "VARCHAR", "CHAR"],
    "DOUBLE PRECISION": ["DOUBLE PRECISION", "FLOAT", "REAL", "NUMERIC", "DECIMAL",
                         "INTEGER", "BIGINT", "SMALLINT", "TEXT", "VARCHAR", "CHAR"],
    "NUMERIC":          ["NUMERIC", "DECIMAL", "REAL", "FLOAT", "DOUBLE PRECISION",
                         "INTEGER", "BIGINT", "SMALLINT", "TEXT", "VARCHAR", "CHAR"],
    "DECIMAL":          ["DECIMAL", "NUMERIC", "REAL", "FLOAT", "DOUBLE PRECISION",
                         "INTEGER", "BIGINT", "SMALLINT", "TEXT", "VARCHAR", "CHAR"],
    # Text family
    "TEXT":             ["TEXT", "VARCHAR", "CHAR"],
    "VARCHAR":          ["VARCHAR", "TEXT", "CHAR"],
    "CHAR":             ["CHAR", "VARCHAR", "TEXT"],
    # Boolean
    "BOOLEAN":          ["BOOLEAN", "INTEGER", "SMALLINT", "TEXT", "VARCHAR", "CHAR"],
    # Date / time family
    "DATE":             ["DATE", "DATETIME", "TIMESTAMP", "TIMESTAMPTZ", "TEXT", "VARCHAR", "CHAR"],
    "DATETIME":         ["DATETIME", "DATE", "TIMESTAMP", "TIMESTAMPTZ", "TEXT", "VARCHAR", "CHAR"],
    "TIMESTAMP":        ["TIMESTAMP", "TIMESTAMPTZ", "DATETIME", "DATE", "TEXT", "VARCHAR", "CHAR"],
    "TIMESTAMPTZ":      ["TIMESTAMPTZ", "TIMESTAMP", "DATETIME", "DATE", "TEXT", "VARCHAR", "CHAR"],
    "TIME":             ["TIME", "TEXT", "VARCHAR", "CHAR"],
    # Identity / opaque — only text representations
    "UUID":             ["UUID", "TEXT", "VARCHAR", "CHAR"],
    # JSON family
    "JSON":             ["JSON", "JSONB", "TEXT", "VARCHAR"],
    "JSONB":            ["JSONB", "JSON", "TEXT", "VARCHAR"],
}


def _are_types_mapping_compatible(origin_type: str, dest_type: str) -> bool:
    """Return True if a column of *origin_type* may be mapped to *dest_type*.

    Both arguments are accepted in any case and may include type parameters
    (e.g. ``"varchar(255)"``); normalisation is delegated to
    :meth:`TypeMapper.extract_type_name`.

    Unknown types fail closed (returns False) to prevent silent data corruption.
    """
    origin_norm = TypeMapper.extract_type_name(origin_type)
    dest_norm = TypeMapper.extract_type_name(dest_type)
    allowed = _TYPE_COMPATIBILITY.get(origin_norm)
    if allowed is None:
        return False
    return dest_norm in allowed


class MappingCreate(BaseModel):
    origin_table: str
    origin_column: str
    destination_table: str
    destination_column: str


def utcnow_iso():
    return datetime.now(timezone.utc).isoformat()


def _plain(value):
    """Los UUID de las consultas SQL directas se devuelven como str (no como uuid.UUID)."""
    return str(value) if isinstance(value, UUID) else value


def _fetch_all(db, sql: str, params: Optional[Dict[str, Any]] = None) -> List[Dict[str, Any]]:
    """Ejecuta SQL con parámetros con nombre (:p) en la Session y devuelve filas como dict."""
    rows = db.execute(sql_text(sql), params or {}).mappings().all()
    return [{k: _plain(v) for k, v in row.items()} for row in rows]


def _fetch_one(db, sql: str, params: Optional[Dict[str, Any]] = None) -> Optional[Dict[str, Any]]:
    """Como _fetch_all pero devuelve la primera fila (o None)."""
    row = db.execute(sql_text(sql), params or {}).mappings().first()
    return {k: _plain(v) for k, v in row.items()} if row else None


def _describe_import_path(data: Any, path) -> str:
    """Ubicación legible (neutral al idioma) de un error de esquema en el JSON importado.

    Ej.: path ('tables', 0, 'columns', 2) -> 'tables[0] ("orders") › columns[2] ("id")'.
    """
    parts: List[str] = []
    node = data
    for key in path:
        try:
            node = node[key]
        except (KeyError, IndexError, TypeError):
            node = None
        if isinstance(key, int):
            label = f"[{key}]"
            name = node.get("name") if isinstance(node, dict) else None
            if isinstance(name, str) and name:
                label += f' ("{name[:60]}")'
            if parts:
                parts[-1] += label
            else:
                parts.append(label)
        else:
            parts.append(str(key))
    return " › ".join(parts) or "$"


def _import_schema_error(ve: schema_exceptions.ValidationError, data: Any) -> Dict[str, Any]:
    """Convierte un ValidationError de jsonschema en un error con código, para que el
    frontend lo traduzca y le diga al usuario qué campo falla y dónde."""
    path = list(ve.absolute_path)

    if ve.validator == "required" and isinstance(ve.instance, dict):
        missing = [f for f in ve.validator_value if f not in ve.instance]
        return make_error(
            ErrorCode.IMPORT_MISSING_FIELD,
            field=", ".join(missing),
            where=_describe_import_path(data, path),
        )

    if ve.validator == "additionalProperties" and isinstance(ve.instance, dict):
        allowed = ve.schema.get("properties", {})
        unknown = [str(k) for k in ve.instance if k not in allowed]
        return make_error(
            ErrorCode.IMPORT_UNKNOWN_FIELD,
            field=", ".join(unknown),
            where=_describe_import_path(data, path),
        )

    if ve.validator == "enum" and path and isinstance(path[-1], str):
        return make_error(
            ErrorCode.IMPORT_INVALID_VALUE,
            field=path[-1],
            where=_describe_import_path(data, path[:-1]),
            value=str(ve.instance)[:60],
            allowed=", ".join(str(v) for v in ve.validator_value),
        )

    return make_error(
        ErrorCode.IMPORT_SCHEMA_INVALID,
        where=_describe_import_path(data, path),
        reason=ve.message[:200],
    )


class Step:
    def __init__(self, name):
        self.name, self.t = name, time.perf_counter()

    def log(self, msg=""):
        dt = (time.perf_counter() - self.t) * 1000
        log.info(f"[import] {self.name} {msg} ({dt:.1f} ms) @ {utcnow_iso()}")
        self.t = time.perf_counter()


@router.get("/projects/{project_id}/export")
def export_mappings(
    request: Request,
    project_id: str,
    format: str = "json",
    db=Depends(get_db_session),
    current_user=Depends(get_current_user),
):
    import sys
    import traceback
    from uuid import UUID

    # Forzar flush de stdout
    sys.stdout.flush()
    sys.stderr.flush()

    log.info(f"[EXPORT] Starting export for project_id={project_id}, format={format}")

    # Validar que project_id es un UUID válido
    try:
        UUID(project_id)
    except ValueError:
        raise HTTPException(
            status_code=400, detail=f"Invalid project_id format. Expected UUID, got: {project_id}"
        )

    try:
        _ = make_translator(request)

        # Validar permisos: owner, editor o viewer
        log.info(f"[EXPORT] Checking permissions...")
        proj = _fetch_one(db, "SELECT owner_id FROM projects WHERE id = :pid", {"pid": project_id})

        if not proj:
            log.error(f"[EXPORT] Project not found")
            raise HTTPException(status_code=404, detail=_("mappings.project_not_found"))

        log.info(f"[EXPORT] Project found, owner_id={proj['owner_id']}")

        if proj["owner_id"] != str(current_user["id"]):
            membership = _fetch_one(
                db,
                """
                SELECT role FROM project_members
                WHERE project_id = :pid AND user_id = :uid
            """,
                {"pid": project_id, "uid": current_user["id"]},
            )
            if not membership:
                raise HTTPException(status_code=403, detail=_("mappings.no_permission"))

            if membership["role"] not in ("viewer", "editor"):
                raise HTTPException(status_code=403, detail=_("mappings.no_permission"))

        # Obtener todos los mappings
        log.info(f"[EXPORT] Getting mappings...")
        # La tabla mappings no tiene columna notes: se devuelve NULL para conservar el formato
        mappings = _fetch_all(
            db,
            """
            SELECT origin_table, origin_column, destination_table, destination_column,
                   NULL AS notes
            FROM mappings WHERE project_id = :pid
        """,
            {"pid": project_id},
        )
        log.info(f"[EXPORT] Found {len(mappings)} mappings")

        if format != "excel" and not mappings:
            log.warning(f"[EXPORT] No mappings found")
            raise HTTPException(status_code=404, detail=_("mappings.no_mappings_found"))

        # Para Excel necesitamos validar columnas obligatorias solo si hay mappings
        if mappings:
            # Agrupar columnas mapeadas por tabla.dest_column
            mapped_columns_by_table: Dict[Any, Set] = {}
            for row in mappings:
                dest_table = row["destination_table"] if isinstance(row, dict) else row[2]
                dest_column = row["destination_column"] if isinstance(row, dict) else row[3]
                mapped_columns_by_table.setdefault(dest_table, set()).add(dest_column)

            # Obtener todas las tablas destino del proyecto
            destination_tables = _fetch_all(
                db,
                """
                SELECT id, name FROM tables
                WHERE project_id = :pid AND type = 'destination'
            """,
                {"pid": project_id},
            )

            for table_row in destination_tables:
                table_id = table_row["id"] if isinstance(table_row, dict) else table_row[0]
                table_name = table_row["name"] if isinstance(table_row, dict) else table_row[1]

                mapped_cols = mapped_columns_by_table.get(table_name)
                if not mapped_cols:
                    continue  # Saltamos comprobación si no hay mapeos para esa tabla

                # Obtener columnas de esa tabla destino
                columns = _fetch_all(
                    db,
                    """
                    SELECT name, is_nullable, is_primary, COALESCE(is_autogenerated, false) AS is_autogenerated
                    FROM columns
                    WHERE table_id = :tid AND table_type = 'destination'
                """,
                    {"tid": table_id},
                )

                for col in columns:
                    name = col["name"] if isinstance(col, dict) else col[0]
                    is_nullable = col["is_nullable"] if isinstance(col, dict) else col[1]
                    is_primary = col["is_primary"] if isinstance(col, dict) else col[2]
                    is_autogenerated = col["is_autogenerated"] if isinstance(col, dict) else col[3]

                    if name not in mapped_cols:
                        if is_autogenerated:
                            continue

                        if not is_nullable:
                            raise HTTPException(
                                status_code=400,
                                detail=_("mappings.not_mapped_not_null").format(
                                    column=name, table=table_name
                                ),
                            )
                        if is_primary:
                            raise HTTPException(
                                status_code=400,
                                detail=_("mappings.not_mapped_primary_key").format(
                                    column=name, table=table_name
                                ),
                            )

        # Si pasa validación, exportamos
        log.info(f"[EXPORT] Starting format export: {format}")
        if format == "json":
            # Obtener información del proyecto
            proj_info = _fetch_one(
                db,
                """
                SELECT name, origin_db, destination_db, created_at
                FROM projects WHERE id = :pid
            """,
                {"pid": project_id},
            )

            if not proj_info:
                raise HTTPException(status_code=404, detail=make_error(ErrorCode.PROJECT_NOT_FOUND))

            # Obtener metadatos de columnas
            columns_info = _fetch_all(
                db,
                """
                SELECT t.name as table_name, t.type as table_type,
                    c.name as column_name, c.data_type, c.is_nullable, c.is_primary
                FROM tables t
                JOIN columns c ON t.id = c.table_id
                WHERE t.project_id = :pid
            """,
                {"pid": project_id},
            )

            # Crear diccionario de metadatos
            metadata = {}
            for col in columns_info:
                table_key = f"{col['table_type']}.{col['table_name']}.{col['column_name']}"
                metadata[table_key] = {
                    "data_type": col["data_type"],
                    "nullable": col["is_nullable"],
                    "primary_key": col["is_primary"],
                }

            # Formato empresarial compatible con herramientas ETL
            project_name = proj_info["name"] or "Untitled Project"
            source_db = proj_info["origin_db"] or "source"
            target_db = proj_info["destination_db"] or "target"

            data = {
                "metadata": {
                    "project_name": project_name,
                    "source_database": source_db,
                    "target_database": target_db,
                    "export_date": datetime.now(timezone.utc).isoformat(),
                    "total_mappings": len(mappings),
                    "format_version": "1.0",
                },
                "mappings": [
                    {
                        "source": {
                            "table": row["origin_table"],
                            "column": row["origin_column"],
                            "data_type": metadata.get(
                                f"origin.{row['origin_table']}.{row['origin_column']}", {}
                            ).get("data_type", "UNKNOWN"),
                            "nullable": metadata.get(
                                f"origin.{row['origin_table']}.{row['origin_column']}", {}
                            ).get("nullable", True),
                            "primary_key": metadata.get(
                                f"origin.{row['origin_table']}.{row['origin_column']}", {}
                            ).get("primary_key", False),
                        },
                        "target": {
                            "table": row["destination_table"],
                            "column": row["destination_column"],
                            "data_type": metadata.get(
                                f"destination.{row['destination_table']}.{row['destination_column']}",
                                {},
                            ).get("data_type", "UNKNOWN"),
                            "nullable": metadata.get(
                                f"destination.{row['destination_table']}.{row['destination_column']}",
                                {},
                            ).get("nullable", True),
                            "primary_key": metadata.get(
                                f"destination.{row['destination_table']}.{row['destination_column']}",
                                {},
                            ).get("primary_key", False),
                        },
                        "transformation": {"type": "direct", "notes": row["notes"]},
                    }
                    for row in mappings
                ],
            }

            buffer = io.BytesIO()
            buffer.write(json.dumps(data, indent=2).encode("utf-8"))
            buffer.seek(0)
            return StreamingResponse(
                buffer,
                media_type="application/json",
                headers={
                    "Content-Disposition": f"attachment; filename=data_mapping_{project_id}.json"
                },
            )

        elif format == "csv":
            # Obtener metadatos para formato empresarial
            columns_info = _fetch_all(
                db,
                """
                SELECT t.name as table_name, t.type as table_type,
                    c.name as column_name, c.data_type, c.is_nullable, c.is_primary
                FROM tables t
                JOIN columns c ON t.id = c.table_id
                WHERE t.project_id = :pid
            """,
                {"pid": project_id},
            )

            metadata = {}
            for col in columns_info:
                table_key = f"{col['table_type']}.{col['table_name']}.{col['column_name']}"
                metadata[table_key] = {
                    "data_type": col["data_type"],
                    "nullable": col["is_nullable"],
                    "primary_key": col["is_primary"],
                }

            csv_buffer = io.StringIO()
            writer = csv.writer(csv_buffer, delimiter=",", quoting=csv.QUOTE_MINIMAL)

            # Header con más información para herramientas ETL
            writer.writerow(
                [
                    "source_table",
                    "source_column",
                    "source_type",
                    "source_nullable",
                    "source_pk",
                    "target_table",
                    "target_column",
                    "target_type",
                    "target_nullable",
                    "target_pk",
                    "transformation_type",
                    "notes",
                ]
            )

            for row in mappings:
                origin_table = row["origin_table"] if isinstance(row, dict) else row[0]
                origin_column = row["origin_column"] if isinstance(row, dict) else row[1]
                dest_table = row["destination_table"] if isinstance(row, dict) else row[2]
                dest_column = row["destination_column"] if isinstance(row, dict) else row[3]

                source_meta = metadata.get(f"origin.{origin_table}.{origin_column}", {})
                target_meta = metadata.get(f"destination.{dest_table}.{dest_column}", {})

                writer.writerow(
                    [
                        origin_table,
                        origin_column,
                        source_meta.get("data_type", "UNKNOWN"),
                        "YES" if source_meta.get("nullable", True) else "NO",
                        "YES" if source_meta.get("primary_key", False) else "NO",
                        dest_table,
                        dest_column,
                        target_meta.get("data_type", "UNKNOWN"),
                        "YES" if target_meta.get("nullable", True) else "NO",
                        "YES" if target_meta.get("primary_key", False) else "NO",
                        "DIRECT",
                        row["notes"] if isinstance(row, dict) else row[4] or "",
                    ]
                )

            byte_stream = io.BytesIO(csv_buffer.getvalue().encode("utf-8-sig"))  # BOM para Excel
            return StreamingResponse(
                byte_stream,
                media_type="text/csv",
                headers={
                    "Content-Disposition": f"attachment; filename=data_mapping_{project_id}.csv"
                },
            )

        elif format == "excel":

            # 1) Traer todas las columnas de tablas ORIGIN
            origin_data = _fetch_all(
                db,
                """
                SELECT t.name as table_name,
                    c.name as column_name,
                    c.data_type,
                    c.is_primary,
                    c.is_nullable
                FROM tables t
                LEFT JOIN columns c ON t.id = c.table_id
                WHERE t.project_id = :pid AND t.type = 'origin'
                ORDER BY t.name, c.id
            """,
                {"pid": project_id},
            )

            # 2) Construir un multimap: origin_table.origin_column -> lista de (dest_table, dest_column)
            mappings_map: dict[str, list[tuple[str, str]]] = {}
            for row in mappings or []:
                origin_table = row["origin_table"] if isinstance(row, dict) else row[0]
                origin_column = row["origin_column"] if isinstance(row, dict) else row[1]
                dest_table = row["destination_table"] if isinstance(row, dict) else row[2]
                dest_column = row["destination_column"] if isinstance(row, dict) else row[3]
                key = f"{origin_table}.{origin_column}"
                mappings_map.setdefault(key, []).append((dest_table or "", dest_column or ""))

            # 3) Organizar datos por tabla origen (con cadenas \n-join para los destinos)
            tables_data: dict[str, list[dict]] = {}
            for row in origin_data:
                table_name = row["table_name"] if isinstance(row, dict) else row[0]
                column_name = row["column_name"] if isinstance(row, dict) else row[1]
                data_type = row["data_type"] if isinstance(row, dict) else row[2]
                is_primary = row["is_primary"] if isinstance(row, dict) else row[3]
                is_nullable = row["is_nullable"] if isinstance(row, dict) else row[4]

                if table_name not in tables_data:
                    tables_data[table_name] = []

                if not column_name:
                    continue

                mapping_key = f"{table_name}.{column_name}"
                pairs = mappings_map.get(mapping_key, [])

                # De-duplicar por si hay repetidos
                pairs = list(dict.fromkeys(pairs))

                dest_tables_str = "\n".join(p[0] for p in pairs) if pairs else ""
                dest_columns_str = "\n".join(p[1] for p in pairs) if pairs else ""

                tables_data[table_name].append(
                    {
                        "column_name": column_name,
                        "data_type": data_type,
                        "is_primary": is_primary,
                        "is_nullable": is_nullable,
                        "dest_table": dest_tables_str,
                        "dest_column": dest_columns_str,
                    }
                )

            # 4) Construcción del Excel
            wb = Workbook()
            wb.remove(wb.active)  # hoja por defecto fuera

            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            center_alignment = Alignment(horizontal="center")
            wrap_alignment = Alignment(wrap_text=True, vertical="top")

            for table_name, columns in tables_data.items():
                if not columns:
                    continue

                ws = wb.create_sheet(title=table_name[:31])

                headers = [
                    "Column Name",
                    "Type",
                    "Primary Key",
                    "Nullable",
                    "Destination Table",
                    "Destination Column",
                ]
                for col_idx, header in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=col_idx, value=header)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_alignment

                for row_idx, col_data in enumerate(columns, 2):
                    ws.cell(row=row_idx, column=1, value=col_data["column_name"])
                    ws.cell(row=row_idx, column=2, value=col_data["data_type"])
                    ws.cell(row=row_idx, column=3, value="Yes" if col_data["is_primary"] else "No")
                    ws.cell(row=row_idx, column=4, value="Yes" if col_data["is_nullable"] else "No")

                    c5 = ws.cell(row=row_idx, column=5, value=col_data["dest_table"])
                    c6 = ws.cell(row=row_idx, column=6, value=col_data["dest_column"])
                    c5.alignment = wrap_alignment
                    c6.alignment = wrap_alignment

                    # Ajuste de altura de fila según nº de líneas
                    lines = max(
                        (col_data["dest_table"].count("\n") + 1) if col_data["dest_table"] else 1,
                        (col_data["dest_column"].count("\n") + 1) if col_data["dest_column"] else 1,
                    )
                    # 15 es una altura razonable por línea; ajústalo si lo ves bajo/alto
                    ws.row_dimensions[row_idx].height = 15 * lines

                # Auto ancho (límite 50)
                for col in ws.columns:
                    max_length = 0
                    column_letter = col[0].column_letter
                    for cell in col:
                        try:
                            max_length = max(
                                max_length, len(str(cell.value)) if cell.value is not None else 0
                            )
                        except:
                            pass
                    ws.column_dimensions[column_letter].width = min(max_length + 2, 50)

            if not wb.sheetnames:
                ws = wb.create_sheet(title="Sin tablas origen")
                ws.cell(row=1, column=1, value="No se encontraron tablas origen en este proyecto")

            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)

            return StreamingResponse(
                buffer,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={
                    "Content-Disposition": f"attachment; filename=mappings_complete_{project_id}.xlsx"
                },
            )

        else:
            raise HTTPException(status_code=400, detail=_("schemas.unsupported_format"))

    except HTTPException as he:
        log.error(f"[EXPORT] HTTPException: {he.status_code} - {he.detail}")
        raise
    except Exception as e:
        error_msg = f"[EXPORT ERROR] {str(e)}\n{traceback.format_exc()}"
        log.error(error_msg)
        raise HTTPException(status_code=500, detail=f"Export error: {str(e)}")


@router.get("/projects/{project_id}/tables")
async def get_project_tables(
    project_id: str,
    request: Request,
    db=Depends(get_db_session),
    current_user=Depends(get_current_user),
):
    """
    Obtiene todas las tablas y columnas de un proyecto
    """
    from app.crud import master_key_crud

    try:
        # Verificar acceso al proyecto
        project = _fetch_one(
            db,
            """
            SELECT owner_id FROM projects WHERE id = :pid
        """,
            {"pid": project_id},
        )

        if not project:
            raise HTTPException(status_code=404, detail=make_error(ErrorCode.PROJECT_NOT_FOUND))

        # Verificar si es owner o miembro
        member = _fetch_one(
            db,
            """
            SELECT role FROM project_members
            WHERE project_id = :pid AND user_id = :uid
        """,
            {"pid": project_id, "uid": current_user["id"]},
        )

        is_owner = project["owner_id"] == str(current_user["id"])
        if not is_owner and not member:
            raise HTTPException(status_code=403, detail=make_error(ErrorCode.PROJECT_NO_ACCESS))

        # Get encryption key from session if available
        encryption_key = getattr(request.state, "encryption_key", None)

        # Debug logging
        import logging

        logger = logging.getLogger(__name__)
        logger.info(f"[GET /tables] project_id: {project_id[:8]}...")
        logger.info(f"[GET /tables] encryption_key available: {encryption_key is not None}")

        #  VALIDAR SESIÓN PENDIENTE
        # El middleware guardó la sesión pero no pudo validarla porque current_user no estaba disponible aún
        if not encryption_key:
            pending_session = getattr(request.state, "_pending_encryption_session", None)
            if pending_session:
                logger.info(f"[GET /tables] Found pending encryption session, validating...")
                # Validar que la sesión pertenece al usuario actual
                session_user_id = pending_session.get("user_id")
                current_user_id = current_user.get("id")

                if session_user_id and current_user_id and session_user_id == current_user_id:
                    #  Sesión válida
                    encryption_key = pending_session["encryption_key"]
                    request.state.encryption_key = encryption_key
                    request.state.encryption_project_id = pending_session["project_id"]
                    logger.info(f"[GET /tables]  Pending session validated and injected")
                else:
                    logger.warning(
                        f"[GET /tables] ⚠️ Pending session validation failed: user mismatch"
                    )
            else:
                logger.info(f"[GET /tables] No pending encryption session found")

        if encryption_key:
            logger.info(f"[GET /tables] encryption_key length: {len(encryption_key)}")

        # Check if project has encryption enabled
        encryption_config = _fetch_one(
            db,
            """
            SELECT encryption_enabled FROM project_encryption
            WHERE project_id = :pid
        """,
            {"pid": project_id},
        )
        is_encrypted = encryption_config and encryption_config.get("encryption_enabled", False)
        logger.info(f"[GET /tables] is_encrypted: {is_encrypted}")

        # Obtener tablas con campos de encriptación si están disponibles
        if is_encrypted:
            tables_raw = _fetch_all(
                db,
                """
                SELECT id, project_id, name, type, created_at, updated_at,
                       name_encrypted, name_nonce
                FROM tables
                WHERE project_id = :pid
                ORDER BY type, name
            """,
                {"pid": project_id},
            )
        else:
            tables_raw = _fetch_all(
                db,
                """
                SELECT id, project_id, name, type, created_at, updated_at
                FROM tables
                WHERE project_id = :pid
                ORDER BY type, name
            """,
                {"pid": project_id},
            )

        # Obtener columnas para cada tabla
        table_ids = [t["id"] for t in tables_raw]
        columns_by_table: Dict[Any, list] = {}

        if table_ids:
            if is_encrypted:
                columns_raw = _fetch_all(
                    db,
                    """
                    SELECT id, table_id, name, data_type, table_type,
                           is_nullable, is_primary, is_autogenerated,
                           created_at, updated_at,
                           name_encrypted, name_nonce,
                           data_type_encrypted, data_type_nonce
                    FROM columns
                    WHERE table_id = ANY(CAST(:ids AS uuid[]))
                    ORDER BY name
                """,
                    {"ids": table_ids},
                )
            else:
                columns_raw = _fetch_all(
                    db,
                    """
                    SELECT id, table_id, name, data_type, table_type,
                           is_nullable, is_primary, is_autogenerated,
                           created_at, updated_at
                    FROM columns
                    WHERE table_id = ANY(CAST(:ids AS uuid[]))
                    ORDER BY name
                """,
                    {"ids": table_ids},
                )

            for col in columns_raw:
                col_dict = dict(col)
                logger.debug(
                    f"[GET /tables] Column raw: name={col_dict.get('name')}, has_encrypted={col_dict.get('name_encrypted') is not None}"
                )

                # Decrypt column fields if encrypted
                if is_encrypted and encryption_key:
                    try:
                        # Decrypt column name
                        if col_dict.get("name_encrypted") and col_dict.get("name_nonce"):
                            decrypted_name = master_key_crud.decrypt_column_name(
                                col_dict["name_encrypted"], col_dict["name_nonce"], encryption_key
                            )
                            col_dict["name"] = decrypted_name
                            logger.info(f"[GET /tables] Decrypted column name: {decrypted_name}")

                        # Decrypt column type
                        if col_dict.get("data_type_encrypted") and col_dict.get("data_type_nonce"):
                            decrypted_type = master_key_crud.decrypt_column_type(
                                col_dict["data_type_encrypted"],
                                col_dict["data_type_nonce"],
                                encryption_key,
                            )
                            col_dict["data_type"] = decrypted_type
                            logger.info(f"[GET /tables] Decrypted column type: {decrypted_type}")
                    except Exception as e:
                        logger.error(
                            f"[GET /tables] Error decrypting column: {str(e)}", exc_info=True
                        )
                        pass
                elif is_encrypted and not encryption_key:
                    logger.warning(
                        f"[GET /tables] Encrypted project but no encryption_key available!"
                    )

                # Remove encrypted fields
                col_dict.pop("name_encrypted", None)
                col_dict.pop("name_nonce", None)
                col_dict.pop("data_type_encrypted", None)
                col_dict.pop("data_type_nonce", None)

                tid = col_dict["table_id"]
                if tid not in columns_by_table:
                    columns_by_table[tid] = []
                columns_by_table[tid].append(col_dict)

        # Construir respuesta con desencriptación de tablas
        result = []
        for table in tables_raw:
            table_dict = dict(table)
            logger.debug(
                f"[GET /tables] Table raw: name={table_dict.get('name')}, has_encrypted={table_dict.get('name_encrypted') is not None}"
            )

            # Decrypt table name if encrypted
            if is_encrypted and encryption_key:
                try:
                    if table_dict.get("name_encrypted") and table_dict.get("name_nonce"):
                        decrypted_name = master_key_crud.decrypt_table_name(
                            table_dict["name_encrypted"], table_dict["name_nonce"], encryption_key
                        )
                        table_dict["name"] = decrypted_name
                        logger.info(f"[GET /tables] Decrypted table name: {decrypted_name}")
                    else:
                        logger.warning(
                            f"[GET /tables] Table {(table_dict.get('id') or '')[:8]}... has no encrypted name fields"
                        )
                except Exception as e:
                    logger.error(
                        f"[GET /tables] Error decrypting table name: {str(e)}", exc_info=True
                    )
                    pass

            # Remove encrypted fields
            table_dict.pop("name_encrypted", None)
            table_dict.pop("name_nonce", None)

            table_dict["columns"] = columns_by_table.get(table_dict["id"], [])
            result.append(table_dict)

        return {"tables": result}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/projects/{project_id}")
async def get_project_mappings(
    project_id: str,
    request: Request,
    db=Depends(get_db_session),
    current_user=Depends(get_current_user),
):
    """
    Obtiene todos los mappings de un proyecto
    """
    from sqlalchemy import text

    from app.crud import mappings_crud, master_key_crud

    try:
        # Verificar acceso
        project = (
            db.execute(text("SELECT owner_id FROM projects WHERE id = :pid"), {"pid": project_id})
            .mappings()
            .fetchone()
        )

        if not project:
            raise HTTPException(status_code=404, detail=make_error(ErrorCode.PROJECT_NOT_FOUND))

        member = (
            db.execute(
                text("SELECT role FROM project_members WHERE project_id = :pid AND user_id = :uid"),
                {"pid": project_id, "uid": current_user["id"]},
            )
            .mappings()
            .fetchone()
        )

        is_owner = str(project["owner_id"]) == str(current_user["id"])
        if not is_owner and not member:
            raise HTTPException(status_code=403, detail=make_error(ErrorCode.PROJECT_NO_ACCESS))

        # Get encryption key from session if available
        encryption_key = getattr(request.state, "encryption_key", None)

        # Debug logging
        import logging

        logger = logging.getLogger(__name__)
        logger.info(f"[GET /mappings] project_id: {project_id[:8]}...")
        logger.info(f"[GET /mappings] encryption_key available: {encryption_key is not None}")

        #  VALIDAR SESIÓN PENDIENTE
        if not encryption_key:
            pending_session = getattr(request.state, "_pending_encryption_session", None)
            if pending_session:
                logger.info(f"[GET /mappings] Found pending encryption session, validating...")
                session_user_id = pending_session.get("user_id")
                current_user_id = current_user.get("id")

                if session_user_id and current_user_id and session_user_id == current_user_id:
                    encryption_key = pending_session["encryption_key"]
                    request.state.encryption_key = encryption_key
                    request.state.encryption_project_id = pending_session["project_id"]
                    logger.info(f"[GET /mappings]  Pending session validated and injected")
                else:
                    logger.warning(f"[GET /mappings] Pending session validation failed")
            else:
                logger.info(f"[GET /mappings] No pending encryption session found")

        if encryption_key:
            logger.info(f"[GET /mappings] encryption_key length: {len(encryption_key)}")

        #  Usar CRUD con encriptación automática
        mappings_raw = mappings_crud.get_mappings_by_project_orm(
            db, project_id, encryption_key=encryption_key
        )

        # Remove encrypted fields from response (keep only decrypted data)
        mappings = []
        for mapping in mappings_raw:
            clean_mapping = {
                "id": mapping.get("id"),
                "project_id": mapping.get("project_id"),
                "origin_table": mapping.get("origin_table"),
                "origin_column": mapping.get("origin_column"),
                "destination_table": mapping.get("destination_table"),
                "destination_column": mapping.get("destination_column"),
                "notes": mapping.get("notes"),
                "created_at": mapping.get("created_at"),
                "updated_at": mapping.get("updated_at"),
            }
            mappings.append(clean_mapping)

        return {"mappings": mappings}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# Añadir un nuevo mapping
@router.post("/projects/{project_id}/save")
async def save_mapping(
    project_id: str,
    mapping: MappingCreate,
    request: Request,
    db_orm=Depends(get_db_session),
    current_user=Depends(get_current_user),
    # NUEVOS parámetros para modo híbrido (query params opcionales)
    origin_source_type: str = "import",
    destination_source_type: str = "import",
    origin_session_id: Optional[str] = None,
    destination_session_id: Optional[str] = None,
    origin_table_id: Optional[str] = None,
    destination_table_id: Optional[str] = None,
):
    current_time = datetime.now(timezone.utc).replace(tzinfo=None)

    try:
        print(f"[MAPPING DEBUG] Received mapping request:")
        print(f"  project_id: {project_id}")
        print(f"  origin_table: '{mapping.origin_table}'")
        print(f"  origin_column: '{mapping.origin_column}'")
        print(f"  destination_table: '{mapping.destination_table}'")
        print(f"  destination_column: '{mapping.destination_column}'")

        _ = make_translator(request)

        # usar flash por redirect
        def redirect_with_flash(msg: str, status_code: int = status.HTTP_303_SEE_OTHER):
            resp = RedirectResponse(url=f"/project_visual/{project_id}", status_code=status_code)
            resp.set_cookie("flash", msg, max_age=4)
            return resp

        # Verificar que el proyecto existe
        project = _fetch_one(
            db_orm,
            """
            SELECT owner_id FROM projects WHERE id = :pid
        """,
            {"pid": project_id},
        )

        if not project:
            raise HTTPException(status_code=404, detail=_("mappings.project_not_found"))

        # Verificar si es owner o editor
        member = _fetch_one(
            db_orm,
            """
            SELECT role FROM project_members
            WHERE project_id = :pid AND user_id = :uid
        """,
            {"pid": project_id, "uid": current_user["id"]},
        )
        es_owner = project["owner_id"] == str(current_user["id"])
        es_editor = member and member["role"] == "editor"

        if not es_owner and not es_editor:
            raise HTTPException(status_code=403, detail=_("mappings.no_permission"))

        # === Límites por plan (primero tablas, luego atributos) ===
        has_pro = bool(
            _fetch_one(
                db_orm,
                """
            SELECT 1
            FROM subscriptions
            WHERE project_id = :pid
              AND LOWER(status) IN ('active','trialing')
              AND plan_id = 'pro'
            LIMIT 1
        """,
                {"pid": project_id},
            )
        )
        MAX_TABLES = 140 if has_pro else 20
        MAX_COLS = 200 if has_pro else 60

        # 1) Tablas totales
        total_tables = int(
            _fetch_one(
                db_orm,
                "SELECT COUNT(*) AS c FROM tables WHERE project_id = :pid",
                {"pid": project_id},
            )["c"]
        )
        if total_tables > MAX_TABLES:
            return redirect_with_flash(_("schemas.max_pro") if has_pro else _("schemas.max_free"))

        # 2) Atributos: ninguna tabla del proyecto debe superar MAX_COLS
        if _fetch_one(
            db_orm,
            """
            SELECT 1
            FROM tables t
            LEFT JOIN columns c ON c.table_id = t.id
            WHERE t.project_id = :pid
            GROUP BY t.id
            HAVING COUNT(c.id) > :max_cols
            LIMIT 1
        """,
            {"pid": project_id, "max_cols": MAX_COLS},
        ):
            return redirect_with_flash(
                _("mappings.max_col_pro") if has_pro else _("mappings.max_col_free")
            )

        #  USAR LA CLAVE DE ENCRIPTACIÓN DEL MIDDLEWARE
        encryption_key = getattr(request.state, "encryption_key", None)

        #  DEBUG: Verificar cookies y estado
        print(f"[MAPPING DEBUG] All cookies: {list(request.cookies.keys())}")
        enc_cookie_name = f"enc_session_{project_id}"
        enc_cookie = request.cookies.get(enc_cookie_name)
        print(f"[MAPPING DEBUG] Cookie '{enc_cookie_name}' exists: {enc_cookie is not None}")
        if enc_cookie:
            print(f"[MAPPING DEBUG] Cookie value (first 20 chars): {enc_cookie[:20]}...")

        #  VALIDAR SESIÓN PENDIENTE
        # El middleware guardó la sesión pero no pudo validarla porque current_user no estaba disponible aún
        if not encryption_key:
            pending_session = getattr(request.state, "_pending_encryption_session", None)
            print(f"[MAPPING DEBUG] Pending session exists: {pending_session is not None}")
            if pending_session:
                print(f"[MAPPING DEBUG] Found pending encryption session, validating...")
                print(
                    f"[MAPPING DEBUG] Session user_id: {pending_session.get('user_id', 'N/A')[:8] if pending_session.get('user_id') else 'N/A'}..."
                )
                print(
                    f"[MAPPING DEBUG] Current user_id: {current_user.get('id', 'N/A')[:8] if current_user.get('id') else 'N/A'}..."
                )
                # Validar que la sesión pertenece al usuario actual
                session_user_id = pending_session.get("user_id")
                current_user_id = current_user.get("id")

                if session_user_id and current_user_id and session_user_id == current_user_id:
                    #  Sesión válida
                    encryption_key = pending_session["encryption_key"]
                    request.state.encryption_key = encryption_key
                    request.state.encryption_project_id = pending_session["project_id"]
                    print(f"[MAPPING DEBUG]  Pending session validated and injected")
                else:
                    print(f"[MAPPING DEBUG] ⚠️ Pending session validation failed: user mismatch")
            else:
                print(f"[MAPPING DEBUG] No pending encryption session found")

        print(
            f"[MAPPING DEBUG] Encryption key from middleware: {len(encryption_key) if encryption_key else 0} bytes"
        )

        # ========== VALIDACIONES MODO HÍBRIDO ==========

        # 1. Validar que si source_type == 'connector', session_id esté presente
        if origin_source_type == "connector" and not origin_session_id:
            raise HTTPException(
                status_code=400,
                detail="origin_session_id is required when origin_source_type is 'connector'",
            )

        if destination_source_type == "connector" and not destination_session_id:
            raise HTTPException(
                status_code=400,
                detail="destination_session_id is required when destination_source_type is 'connector'",
            )

        # 2. Validar sesiones si están en modo connector
        from app.services.db_connection_service import (
            DBConnectionService,
            SessionExpiredError,
            SessionNotFoundError,
        )

        if origin_source_type == "connector" and origin_session_id:
            # Obtener servicio de conexiones
            redis_client = get_redis_client()
            encryption_key = os.getenv("ENCRYPTION_MASTER_KEY")
            if not encryption_key:
                raise HTTPException(status_code=500, detail="Encryption master key not configured")
            db_service = DBConnectionService(redis_client, encryption_key)

            try:
                # Validar que la sesión existe y no ha expirado
                session = db_service.get_session(current_user["id"], origin_session_id)

                # Validar ownership (la sesión debe pertenecer al usuario actual)
                if session.user_id != current_user["id"]:
                    raise HTTPException(
                        status_code=403, detail="origin_session_id does not belong to current user"
                    )

            except SessionNotFoundError:
                raise HTTPException(
                    status_code=404, detail=f"origin_session_id '{origin_session_id}' not found"
                )
            except SessionExpiredError:
                raise HTTPException(
                    status_code=410,
                    detail=f"origin_session_id '{origin_session_id}' has expired. Please reconnect to the database.",
                )

        if destination_source_type == "connector" and destination_session_id:
            redis_client = get_redis_client()
            encryption_key = os.getenv("ENCRYPTION_MASTER_KEY")
            if not encryption_key:
                raise HTTPException(status_code=500, detail="Encryption master key not configured")
            db_service = DBConnectionService(redis_client, encryption_key)

            try:
                session = db_service.get_session(current_user["id"], destination_session_id)

                if session.user_id != current_user["id"]:
                    raise HTTPException(
                        status_code=403,
                        detail="destination_session_id does not belong to current user",
                    )

            except SessionNotFoundError:
                raise HTTPException(
                    status_code=404,
                    detail=f"destination_session_id '{destination_session_id}' not found",
                )
            except SessionExpiredError:
                raise HTTPException(
                    status_code=410,
                    detail=f"destination_session_id '{destination_session_id}' has expired. Please reconnect to the database.",
                )

        # 3. Validar table_ids si están en modo import (opcional pero recomendado)
        if origin_source_type == "import" and origin_table_id:
            if not _fetch_one(
                db_orm,
                """
                SELECT id FROM tables
                WHERE id = :tid AND project_id = :pid
            """,
                {"tid": origin_table_id, "pid": project_id},
            ):
                raise HTTPException(
                    status_code=404,
                    detail=f"origin_table_id '{origin_table_id}' not found in project tables",
                )

        if destination_source_type == "import" and destination_table_id:
            if not _fetch_one(
                db_orm,
                """
                SELECT id FROM tables
                WHERE id = :tid AND project_id = :pid
            """,
                {"tid": destination_table_id, "pid": project_id},
            ):
                raise HTTPException(
                    status_code=404,
                    detail=f"destination_table_id '{destination_table_id}' not found in project tables",
                )

        # Usar el CRUD que maneja encriptación automáticamente
        from app.crud import mappings_crud, schemas_crud

        # Buscar tabla origen
        print(f"[MAPPING DEBUG] Searching for origin table: '{mapping.origin_table}'")
        origin_tables = schemas_crud.get_tables_by_project_orm(
            db_orm, project_id, table_type="origin", encryption_key=encryption_key
        )
        origin_table_obj = next(
            (t for t in origin_tables if t.get("name") == mapping.origin_table), None
        )

        if not origin_table_obj:
            print(f"[MAPPING DEBUG] Origin table not found. Available tables:")
            for t in origin_tables:
                print(f"  - {t.get('name')} (id: {t.get('id')})")
            raise HTTPException(status_code=400, detail=_("mappings.origin_not_found"))

        # Buscar columna origen
        print(
            f"[MAPPING DEBUG] Searching for origin column: '{mapping.origin_column}' in table {origin_table_obj['id']}"
        )
        origin_columns = schemas_crud.get_columns_by_table_orm(
            db_orm, origin_table_obj["id"], encryption_key=encryption_key
        )
        origin_col = next(
            (c for c in origin_columns if c.get("name") == mapping.origin_column), None
        )

        print(f"[MAPPING DEBUG] Origin column found: {origin_col is not None}")
        if not origin_col:
            print(f"[MAPPING DEBUG] Available columns in table:")
            for c in origin_columns:
                print(f"  - {c.get('name')}")
            raise HTTPException(status_code=400, detail=_("mappings.origin_not_found"))
        origin_type = origin_col["data_type"]

        # Buscar tabla destino
        print(f"[MAPPING DEBUG] Searching for destination table: '{mapping.destination_table}'")
        destination_tables = schemas_crud.get_tables_by_project_orm(
            db_orm, project_id, table_type="destination", encryption_key=encryption_key
        )
        dest_table_obj = next(
            (t for t in destination_tables if t.get("name") == mapping.destination_table), None
        )

        if not dest_table_obj:
            print(f"[MAPPING DEBUG] Destination table not found. Available tables:")
            for t in destination_tables:
                print(f"  - {t.get('name')} (id: {t.get('id')})")
            raise HTTPException(status_code=400, detail=_("mappings.destination_not_found"))

        # Buscar columna destino
        print(
            f"[MAPPING DEBUG] Searching for destination column: '{mapping.destination_column}' in table {dest_table_obj['id']}"
        )
        dest_columns = schemas_crud.get_columns_by_table_orm(
            db_orm, dest_table_obj["id"], encryption_key=encryption_key
        )
        dest_col = next(
            (c for c in dest_columns if c.get("name") == mapping.destination_column), None
        )

        print(f"[MAPPING DEBUG] Destination column found: {dest_col is not None}")
        if not dest_col:
            print(f"[MAPPING DEBUG] Available columns in table:")
            for c in dest_columns:
                print(f"  - {c.get('name')}")
            raise HTTPException(status_code=400, detail=_("mappings.destination_not_found"))
        dest_type = dest_col["data_type"]

        # --- Validación de compatibilidad de tipos ---
        if not _are_types_mapping_compatible(origin_type, dest_type):
            raise HTTPException(
                status_code=400,
                detail=_("mappings.type_mismatch").format(origin=origin_type, dest=dest_type),
            )

        # Verificar duplicado usando el CRUD con encriptación
        existing_mappings = mappings_crud.get_mappings_by_project_orm(
            db_orm, project_id, encryption_key=encryption_key
        )
        duplicate = next(
            (
                m
                for m in existing_mappings
                if m.get("origin_table") == mapping.origin_table
                and m.get("origin_column") == mapping.origin_column
                and m.get("destination_table") == mapping.destination_table
                and m.get("destination_column") == mapping.destination_column
            ),
            None,
        )
        if duplicate:
            raise HTTPException(status_code=400, detail=_("mappings.duplicate"))

        #  Crear mapping usando CRUD con encriptación automática + parámetros híbridos
        new_mapping = mappings_crud.create_mapping_orm(
            db=db_orm,
            project_id=project_id,
            origin_table=mapping.origin_table,
            origin_column=mapping.origin_column,
            destination_table=mapping.destination_table,
            destination_column=mapping.destination_column,
            user_id=current_user["id"],
            encryption_key=encryption_key,
            # NUEVOS parámetros para modo híbrido
            origin_source_type=origin_source_type,
            destination_source_type=destination_source_type,
            origin_session_id=origin_session_id,
            destination_session_id=destination_session_id,
            origin_table_id=origin_table_id,
            destination_table_id=destination_table_id,
        )

        # CREATE VERSION ENTRY
        try:
            from app.services import version_control_service

            print(f"[VERSION CONTROL] Starting version creation for mapping creation")
            print(f"[VERSION CONTROL] Project ID: {project_id}")
            print(f"[VERSION CONTROL] User ID: {current_user['id']}")
            print(
                f"[VERSION CONTROL] Mapping: {mapping.origin_table}.{mapping.origin_column} → {mapping.destination_table}.{mapping.destination_column}"
            )

            # Get all mappings for snapshot using CRUD (with decryption)
            all_mappings_full = mappings_crud.get_mappings_by_project_orm(
                db_orm, project_id, encryption_key=encryption_key
            )
            all_mappings = [
                {
                    "origin_table": m.get("origin_table"),
                    "origin_column": m.get("origin_column"),
                    "destination_table": m.get("destination_table"),
                    "destination_column": m.get("destination_column"),
                }
                for m in all_mappings_full
            ]
            print(f"[VERSION CONTROL] Total mappings in project: {len(all_mappings)}")

            # Create version snapshot
            snapshot = {
                "mappings": all_mappings,
                "mapping_added": {
                    "origin_table": mapping.origin_table,
                    "origin_column": mapping.origin_column,
                    "destination_table": mapping.destination_table,
                    "destination_column": mapping.destination_column,
                },
            }
            print(f"[VERSION CONTROL] Snapshot size: {len(str(snapshot))} characters")

            version = version_control_service.create_version(
                db=db_orm,
                project_id=project_id,
                entity_id=project_id,  # Use project_id as entity_id for project-level changes
                entity_type="mapping",
                user_id=current_user["id"],
                change_type="create",
                snapshot=snapshot,
                message=f"Mapping created: {mapping.origin_table}.{mapping.origin_column} → {mapping.destination_table}.{mapping.destination_column}",
            )
            print(
                f"[VERSION CONTROL] ✓ Version created successfully: {version.get('id', 'unknown')} - Version #{version.get('version_number', 'unknown')}"
            )
        except Exception as e:
            # Don't fail the mapping creation if version creation fails
            print(f"[VERSION CONTROL] ✗ FAILED to create version for mapping creation")
            print(f"[VERSION CONTROL] Error type: {type(e).__name__}")
            print(f"[VERSION CONTROL] Error message: {str(e)}")
            import traceback

            print(f"[VERSION CONTROL] Traceback:\n{traceback.format_exc()}")

        # EMITIR EVENTO EN TIEMPO REAL
        await realtime_broadcast(
            project_id,
            {
                "type": "mapping_added",
                "ts": datetime.now(timezone.utc).isoformat(),
                "actor_id": str(current_user["id"]),
                "actor_name": current_user.get("username", "Usuario"),
                "mapping": {
                    "project_id": project_id,
                    "origin_table": mapping.origin_table,
                    "origin_column": mapping.origin_column,
                    "destination_table": mapping.destination_table,
                    "destination_column": mapping.destination_column,
                },
            },
        )

        return {"status": "ok"}

    except HTTPException:
        raise
    except Exception as e:
        try:
            db_orm.rollback()
        except Exception:
            pass
        # Log detailed error for debugging
        import traceback

        print(f"[MAPPING ERROR] Exception type: {type(e).__name__}")
        print(f"[MAPPING ERROR] Exception message: {str(e)}")
        print(f"[MAPPING ERROR] Traceback:\n{traceback.format_exc()}")
        raise HTTPException(status_code=500, detail=_("mappings.save_error_generic"))


# Validar mappings del proyecto
@router.post("/projects/{project_id}/validate")
async def validate_mappings(
    project_id: str, db=Depends(get_db_session), current_user=Depends(get_current_user)
):
    """Valida que todos los mappings del proyecto sean correctos"""
    try:
        # Verificar acceso
        project = _fetch_one(
            db, "SELECT owner_id FROM projects WHERE id = :pid", {"pid": project_id}
        )

        if not project:
            raise HTTPException(status_code=404, detail=make_error(ErrorCode.PROJECT_NOT_FOUND))

        member = _fetch_one(
            db,
            """
            SELECT role FROM project_members
            WHERE project_id = :pid AND user_id = :uid
        """,
            {"pid": project_id, "uid": current_user["id"]},
        )

        is_owner = project["owner_id"] == str(current_user["id"])
        if not is_owner and not member:
            raise HTTPException(status_code=403, detail=make_error(ErrorCode.PROJECT_NO_ACCESS))

        errors = []

        # Obtener tablas destino con columnas obligatorias
        required_columns = _fetch_all(
            db,
            """
            SELECT t.name as table_name, c.name as column_name, c.is_nullable, c.is_primary,
                   COALESCE(c.is_autogenerated, false) as is_autogenerated
            FROM tables t
            JOIN columns c ON t.id = c.table_id
            WHERE t.project_id = :pid AND t.type = 'destination'
            AND (c.is_primary = true OR c.is_nullable = false)
        """,
            {"pid": project_id},
        )

        # Obtener mappings existentes
        existing_mappings = _fetch_all(
            db,
            """
            SELECT destination_table, destination_column
            FROM mappings
            WHERE project_id = :pid
        """,
            {"pid": project_id},
        )

        mapped_columns = set()
        for m in existing_mappings:
            mapped_columns.add(f"{m['destination_table']}.{m['destination_column']}")

        # Validar columnas obligatorias
        for col in required_columns:
            if col["is_autogenerated"]:
                continue

            key = f"{col['table_name']}.{col['column_name']}"
            if key not in mapped_columns:
                if col["is_primary"]:
                    errors.append(
                        f"Columna PK no mapeada: {col['table_name']}.{col['column_name']}"
                    )
                elif not col["is_nullable"]:
                    errors.append(
                        f"Columna NOT NULL no mapeada: {col['table_name']}.{col['column_name']}"
                    )

        return {"isValid": len(errors) == 0, "errors": errors}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# Eliminar mapping
@router.delete("/projects/{project_id}/delete")
async def delete_mapping(
    request: Request,
    project_id: str,
    origin_table: str = Body(...),
    origin_column: str = Body(...),
    destination_table: str = Body(...),
    destination_column: str = Body(...),
    db_orm=Depends(get_db_session),
    current_user=Depends(get_current_user),
):
    from app.crud import mappings_crud

    try:
        _ = make_translator(request)

        def redirect_with_flash(msg: str, status_code: int = status.HTTP_303_SEE_OTHER):
            resp = RedirectResponse(url=f"/project_visual/{project_id}", status_code=status_code)
            resp.set_cookie("flash", msg, max_age=4)
            return resp

        # Permisos
        project = _fetch_one(
            db_orm, "SELECT owner_id FROM projects WHERE id = :pid", {"pid": project_id}
        )
        if not project:
            raise HTTPException(status_code=404, detail=_("mappings.project_not_found"))

        if project["owner_id"] != str(current_user["id"]):
            membership = _fetch_one(
                db_orm,
                """
                SELECT role
                FROM project_members
                WHERE project_id = :pid AND user_id = :uid
            """,
                {"pid": project_id, "uid": current_user["id"]},
            )
            if not membership or membership["role"] != "editor":
                # (opcional) volver al proyecto en lugar de /projects
                return redirect_with_flash(_("mappings.no_permission"))

        # --- Límites de plan (mismas que save_mapping) ---
        has_pro = bool(
            _fetch_one(
                db_orm,
                """
            SELECT 1
            FROM subscriptions
            WHERE project_id = :pid
              AND LOWER(status) IN ('active','trialing')
              AND plan_id = 'pro'
            LIMIT 1
        """,
                {"pid": project_id},
            )
        )
        MAX_TABLES = 140 if has_pro else 20
        MAX_COLS = 200 if has_pro else 60

        # 1) Tablas totales
        total_tables = int(
            _fetch_one(
                db_orm,
                "SELECT COUNT(*) AS c FROM tables WHERE project_id = :pid",
                {"pid": project_id},
            )["c"]
        )
        if total_tables > MAX_TABLES:
            return redirect_with_flash(_("schemas.max_pro") if has_pro else _("schemas.max_free"))

        # 2) Atributos por tabla
        if _fetch_one(
            db_orm,
            """
            SELECT 1
            FROM tables t
            LEFT JOIN columns c ON c.table_id = t.id
            WHERE t.project_id = :pid
            GROUP BY t.id
            HAVING COUNT(c.id) > :max_cols
            LIMIT 1
        """,
            {"pid": project_id, "max_cols": MAX_COLS},
        ):
            return redirect_with_flash(
                _("mappings.max_col_pro") if has_pro else _("mappings.max_col_free")
            )

        #  Obtener encryption key del middleware
        encryption_key = getattr(request.state, "encryption_key", None)

        #  VALIDAR SESIÓN PENDIENTE
        if not encryption_key:
            pending_session = getattr(request.state, "_pending_encryption_session", None)
            if pending_session:
                session_user_id = pending_session.get("user_id")
                current_user_id = current_user.get("id")
                if session_user_id and current_user_id and session_user_id == current_user_id:
                    encryption_key = pending_session["encryption_key"]
                    request.state.encryption_key = encryption_key
                    request.state.encryption_project_id = pending_session["project_id"]

        #  Buscar el mapping a eliminar usando CRUD (con desencriptación)
        all_mappings = mappings_crud.get_mappings_by_project_orm(
            db_orm, project_id, encryption_key=encryption_key
        )
        mapping_to_delete = next(
            (
                m
                for m in all_mappings
                if m.get("origin_table") == origin_table
                and m.get("origin_column") == origin_column
                and m.get("destination_table") == destination_table
                and m.get("destination_column") == destination_column
            ),
            None,
        )

        if not mapping_to_delete:
            raise HTTPException(status_code=404, detail=_("mappings.not_found"))

        #  Eliminar mapping usando CRUD
        mappings_crud.delete_mapping_orm(
            db=db_orm, mapping_id=mapping_to_delete["id"], user_id=current_user["id"]
        )

        # CREATE VERSION ENTRY
        try:
            from app.services import version_control_service

            print(f"[VERSION CONTROL] Starting version creation for mapping deletion")
            print(f"[VERSION CONTROL] Project ID: {project_id}")
            print(f"[VERSION CONTROL] User ID: {current_user['id']}")
            print(
                f"[VERSION CONTROL] Mapping: {origin_table}.{origin_column} → {destination_table}.{destination_column}"
            )

            # Get all remaining mappings for snapshot using CRUD (with decryption)
            all_mappings_full = mappings_crud.get_mappings_by_project_orm(
                db_orm, project_id, encryption_key=encryption_key
            )
            all_mappings = [
                {
                    "origin_table": m.get("origin_table"),
                    "origin_column": m.get("origin_column"),
                    "destination_table": m.get("destination_table"),
                    "destination_column": m.get("destination_column"),
                }
                for m in all_mappings_full
            ]
            print(f"[VERSION CONTROL] Remaining mappings in project: {len(all_mappings)}")

            # Create version snapshot
            snapshot = {
                "mappings": all_mappings,
                "mapping_deleted": {
                    "origin_table": origin_table,
                    "origin_column": origin_column,
                    "destination_table": destination_table,
                    "destination_column": destination_column,
                },
            }
            print(f"[VERSION CONTROL] Snapshot size: {len(str(snapshot))} characters")

            version = version_control_service.create_version(
                db=db_orm,
                project_id=project_id,
                entity_id=project_id,  # Use project_id as entity_id for project-level changes
                entity_type="mapping",
                user_id=current_user["id"],
                change_type="delete",
                snapshot=snapshot,
                message=f"Mapping deleted: {origin_table}.{origin_column} → {destination_table}.{destination_column}",
            )
            print(
                f"[VERSION CONTROL] ✓ Version created successfully: {version.get('id', 'unknown')} - Version #{version.get('version_number', 'unknown')}"
            )
        except Exception as e:
            # Don't fail the mapping deletion if version creation fails
            print(f"[VERSION CONTROL] ✗ FAILED to create version for mapping deletion")
            print(f"[VERSION CONTROL] Error type: {type(e).__name__}")
            print(f"[VERSION CONTROL] Error message: {str(e)}")
            import traceback

            print(f"[VERSION CONTROL] Traceback:\n{traceback.format_exc()}")

        await realtime_broadcast(
            project_id,
            {
                "type": "mapping_deleted",
                "ts": datetime.utcnow().isoformat(),
                "actor_id": str(current_user["id"]),
                "actor_name": current_user.get("username", "Usuario"),
                "mapping": {
                    "project_id": project_id,
                    "origin_table": origin_table,
                    "origin_column": origin_column,
                    "destination_table": destination_table,
                    "destination_column": destination_column,
                },
            },
        )

        # mostrar banner inmediatamente al volver
        return {"ok": True, "message": _("mappings.deleted_success")}

    except HTTPException:
        raise
    except Exception as e:
        try:
            db_orm.rollback()
        except Exception:
            pass
        raise HTTPException(status_code=500, detail=_("mappings.save_error_generic"))


# Importar estructura de tablas (JSON + CSV + XLSX) con límites Free/Pro
@router.post("/projects/{project_id}/import_structure")
async def import_structure(
    request: Request,
    project_id: str,
    file: UploadFile = File(...),
    format: str = Form("json"),
    duplicate_strategy: str = Form("skip"),
    db=Depends(get_db_session),
    current_user=Depends(get_current_user),
):
    _ = make_translator(request)

    def ensure_editor_permissions(db):
        proj = _fetch_one(db, "SELECT owner_id FROM projects WHERE id = :pid", {"pid": project_id})
        if not proj:
            raise HTTPException(status_code=404, detail=_("mappings.project_not_found"))
        if proj["owner_id"] == str(current_user["id"]):
            return
        row = _fetch_one(
            db,
            "SELECT role FROM project_members WHERE project_id = :pid AND user_id = :uid",
            {"pid": project_id, "uid": current_user["id"]},
        )
        if not row or row["role"] != "editor":
            raise HTTPException(
                status_code=403,
                detail=_("mappings.no_permission"),
            )

    def redirect_with_flash(msg: str):
        resp = RedirectResponse(
            url=f"/project_visual/{project_id}", status_code=status.HTTP_303_SEE_OTHER
        )
        resp.set_cookie("flash", msg, max_age=4)
        return resp

    ALLOWED_TYPES = {"TEXT", "INTEGER", "REAL", "BOOLEAN", "DATE", "DATETIME"}
    TRUES = {"1", "true", "yes", "y", "t", "si", "sí"}

    current_time = datetime.now(timezone.utc).replace(tzinfo=None)

    # Import CRUD para encriptación
    from app.crud import schemas_crud

    try:
        # Permisos
        ensure_editor_permissions(db)

        #  Obtener encryption key del middleware
        encryption_key = getattr(request.state, "encryption_key", None)

        # ✅ INICIALIZAR DEBUG LOGGER TEMPRANO
        import logging

        debug_logger = logging.getLogger(__name__)

        #  VALIDAR SESIÓN PENDIENTE
        if not encryption_key:
            debug_logger.info(
                "[IMPORT_DEBUG] encryption_key is None, checking for pending session..."
            )
            pending_session = getattr(request.state, "_pending_encryption_session", None)
            if pending_session:
                debug_logger.info("[IMPORT_DEBUG] Pending session found, validating...")
                session_user_id = pending_session.get("user_id")
                current_user_id = current_user.get("id")

                debug_logger.info(
                    f"[IMPORT_DEBUG] User ID comparison:\n"
                    f"  session_user_id: {session_user_id}\n"
                    f"  current_user_id: {current_user_id}\n"
                    f"  match: {session_user_id == current_user_id}"
                )

                if session_user_id and current_user_id and session_user_id == current_user_id:
                    debug_logger.info(
                        "[IMPORT_DEBUG] ✓ User IDs match! Injecting encryption_key..."
                    )
                    encryption_key = pending_session["encryption_key"]
                    request.state.encryption_key = encryption_key
                    request.state.encryption_project_id = pending_session["project_id"]
                    debug_logger.info(
                        f"[IMPORT_DEBUG] ✓ encryption_key injected (length: {len(encryption_key) if encryption_key else 0} bytes)"
                    )
                else:
                    debug_logger.warning(
                        f"[IMPORT_DEBUG] ✗ User ID validation FAILED:\n"
                        f"  session_user_id present: {session_user_id is not None}\n"
                        f"  current_user_id present: {current_user_id is not None}\n"
                        f"  IDs match: {session_user_id == current_user_id if (session_user_id and current_user_id) else 'N/A'}"
                    )
            else:
                debug_logger.warning("[IMPORT_DEBUG] ✗ No pending session found in request.state")

        # ✅ DEBUG LOGGING: Para diagnosticar problemas de sesión de encriptación
        enc_cookie_name = f"enc_session_{project_id}"
        enc_cookie_exists = enc_cookie_name in request.cookies
        pending_session_exists = hasattr(request.state, "_pending_encryption_session")

        # 🔍 EXTENDED DEBUG: Log más detallado
        debug_logger.info(
            f"[IMPORT_DEBUG] ===== ENCRYPTION SESSION DEBUG =====\n"
            f"  project_id: {project_id}\n"
            f"  current_user_id: {current_user.get('id')}\n"
            f"  cookies_available: {list(request.cookies.keys())}\n"
            f"  enc_cookie_exists: {enc_cookie_exists}\n"
            f"  encryption_key_from_middleware: {encryption_key is not None}\n"
            f"  pending_session_exists: {pending_session_exists}\n"
            f"  request.state.encryption_key: {getattr(request.state, 'encryption_key', 'NOT_SET')}\n"
            f"  request.state.encryption_project_id: {getattr(request.state, 'encryption_project_id', 'NOT_SET')}\n"
            f"  request.state.current_user: {getattr(request.state, 'current_user', 'NOT_SET') is not None}"
        )

        # 🔍 Si hay sesión pendiente, loggear detalles
        if pending_session_exists:
            pending_session = getattr(request.state, "_pending_encryption_session", None)
            if pending_session:
                debug_logger.info(
                    f"[IMPORT_DEBUG] Pending session details:\n"
                    f"  session_project_id: {pending_session.get('project_id', 'N/A')}\n"
                    f"  session_user_id: {pending_session.get('user_id', 'N/A')}\n"
                    f"  session_has_encryption_key: {pending_session.get('encryption_key') is not None}\n"
                    f"  session_expires_at: {pending_session.get('expires_at', 'N/A')}"
                )

        # ✅ VALIDACIÓN PRE-VUELO: Verificar encriptación ANTES de procesar archivo
        from app.crud import master_key_crud

        config = master_key_crud.get_project_encryption_orm(db, project_id)

        if config and config["encryption_enabled"]:
            debug_logger.info(f"[IMPORT_DEBUG] Proyecto {project_id} tiene encriptación habilitada")
            if not encryption_key:
                debug_logger.warning(
                    f"[IMPORT_DEBUG] Rechazando import: proyecto requiere encriptación pero encryption_key=None. "
                    f"Cookie existe: {enc_cookie_exists}, Pending session: {pending_session_exists}"
                )
                raise HTTPException(
                    status_code=403,
                    detail=_(
                        "El proyecto requiere encriptación. Por favor, desbloquea el proyecto antes de importar."
                    ),
                )
            else:
                debug_logger.info(
                    f"[IMPORT_DEBUG] encryption_key disponible, procediendo con importación"
                )
        else:
            debug_logger.info(
                f"[IMPORT_DEBUG] Proyecto {project_id} NO tiene encriptación habilitada o no existe config"
            )

        # ========== VALIDACIONES DE SEGURIDAD ==========

        # 1. VALIDACIÓN DE TAMAÑO DE ARCHIVO (BACKEND)
        MAX_FILE_SIZE = 15 * 1024 * 1024  # 15MB
        content = await file.read()
        if len(content) > MAX_FILE_SIZE:
            raise HTTPException(
                status_code=413,
                detail=f"Archivo demasiado grande. Máximo permitido: {MAX_FILE_SIZE // (1024*1024)}MB",
            )

        if not content:
            raise HTTPException(status_code=400, detail=_("schemas.file_empty"))

        # 2. VALIDACIÓN DE NOMBRE DE ARCHIVO
        if not file.filename:
            raise HTTPException(status_code=400, detail=make_error(ErrorCode.FILENAME_REQUIRED))

        # Sanitizar nombre de archivo
        import os
        import re

        safe_filename = os.path.basename(file.filename)
        safe_filename = re.sub(r"[^a-zA-Z0-9._-]", "", safe_filename)

        # 3. VALIDACIÓN DE EXTENSIÓN Y TIPO MIME
        allowed_extensions = {"json": [".json"], "csv": [".csv"], "excel": [".xlsx", ".xls"]}

        allowed_mimes = {
            "json": ["application/json", "text/json"],
            "csv": ["text/csv", "application/csv", "text/plain"],
            "excel": [
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "application/vnd.ms-excel",
            ],
        }

        # Verificar extensión
        file_ext = os.path.splitext(safe_filename)[1].lower()
        if format not in allowed_extensions:
            raise HTTPException(status_code=400, detail=make_error(ErrorCode.UNSUPPORTED_FORMAT))

        if file_ext not in allowed_extensions[format]:
            raise HTTPException(
                status_code=400,
                detail=f"Extensión de archivo no válida para formato {format}. Esperado: {allowed_extensions[format]}",
            )

        # Verificar tipo MIME
        content_type = file.content_type or ""
        if content_type not in allowed_mimes[format]:
            raise HTTPException(
                status_code=400,
                detail=f"Tipo MIME no válido: {content_type}. Esperado: {allowed_mimes[format]}",
            )

        # 4. VALIDACIÓN ADICIONAL DE CONTENIDO
        # Verificar que el contenido no contenga caracteres de control peligrosos
        try:
            # Para archivos de texto, verificar encoding seguro
            if format in ["json", "csv"]:
                test_decode = content.decode("utf-8", errors="strict")
                # Verificar que no contenga caracteres de control peligrosos
                import string

                printable = set(string.printable)
                if not all(
                    c in printable or c.isspace() for c in test_decode[:1000]
                ):  # Verificar primeros 1000 chars
                    raise HTTPException(
                        status_code=400, detail="Archivo contiene caracteres no válidos"
                    )
        except UnicodeDecodeError:
            if format in ["json", "csv"]:
                raise HTTPException(
                    status_code=400, detail="Archivo no tiene codificación UTF-8 válida"
                )

        # --------- Plan (Free / Pro) ----------
        has_pro = bool(
            _fetch_one(
                db,
                """
            SELECT 1
            FROM subscriptions
            WHERE project_id = :pid
              AND LOWER(status) IN ('active','trialing')
              AND plan_id = 'pro'
            LIMIT 1
        """,
                {"pid": project_id},
            )
        )
        max_tables = 140 if has_pro else 20
        MAX_COLS = 200 if has_pro else 60

        # -----------------------------
        #     PARSEO SEGÚN FORMATO
        # -----------------------------
        tables_payload: List[Dict[str, Any]] = []  # [{name, type, columns:[{...}]}]

        if format == "json":
            try:
                data = json.loads(content)
            except Exception:
                raise HTTPException(status_code=400, detail=_("schemas.invalid_json"))
            if (
                not isinstance(data, dict)
                or "tables" not in data
                or not isinstance(data["tables"], list)
            ):
                raise HTTPException(
                    status_code=400,
                    detail=_("schemas.missing_tables_key"),
                )

            # Validación con tu JSONSchema
            schema_path = os.path.join(
                os.path.dirname(__file__), "..", "static", "schemas", "table_import_schema.json"
            )
            try:
                with open(schema_path, "r", encoding="utf-8") as f:
                    schema = json.load(f)
                Draft202012Validator(schema).validate(data)
            except schema_exceptions.ValidationError as ve:
                raise HTTPException(status_code=400, detail=_import_schema_error(ve, data))
            except FileNotFoundError:
                print("[IMPORT][WARN] Schema JSON no encontrado en:", schema_path)

            for tbl in data["tables"]:
                name = str(tbl.get("name", "")).strip()
                ttype = str(tbl.get("type", "")).strip().lower()
                cols = tbl.get("columns") or []
                if not name or ttype not in ("origin", "destination") or not isinstance(cols, list):
                    continue
                norm_cols = []
                for c in cols:
                    cname = str(c.get("name", "")).strip()
                    ctype = str(c.get("type", "")).strip().upper()
                    if not cname or ctype not in ALLOWED_TYPES:
                        continue
                    norm_cols.append(
                        {
                            "name": cname,
                            "type": ctype,
                            "is_primary": bool(c.get("is_primary", False)),
                            "is_nullable": bool(c.get("is_nullable", True)),
                        }
                    )
                tables_payload.append({"name": name, "type": ttype, "columns": norm_cols})

        elif format == "csv":
            try:
                text = content.decode("utf-8", errors="strict")
            except Exception:
                text = content.decode("latin-1", errors="ignore")

            reader = csv.DictReader(StringIO(text))
            required = {
                "table_name",
                "type",
                "column_name",
                "data_type",
                "is_primary",
                "is_nullable",
            }
            if not required.issubset(set(h.strip().lower() for h in (reader.fieldnames or []))):
                raise HTTPException(
                    status_code=400,
                    detail=_("schemas.invalid_csv_headers"),
                )

            tmp = _dd(list)
            for row in reader:
                tname = str(row.get("table_name", "")).strip()
                ttype = str(row.get("type", "")).strip().lower()
                cname = str(row.get("column_name", "")).strip()
                ctype = str(row.get("data_type", "")).strip().upper()
                ip = str(row.get("is_primary", "")).strip().lower()
                inull = str(row.get("is_nullable", "")).strip().lower()

                if not tname or ttype not in ("origin", "destination"):
                    continue
                if not cname or ctype not in ALLOWED_TYPES:
                    continue

                tmp[(tname, ttype)].append(
                    {
                        "name": cname,
                        "type": ctype,
                        "is_primary": ip in TRUES,
                        "is_nullable": inull in TRUES,
                    }
                )

            for (tname, ttype), cols in tmp.items():
                tables_payload.append({"name": tname, "type": ttype, "columns": cols})

        elif format == "excel":
            try:
                from openpyxl import load_workbook
            except ImportError:
                raise HTTPException(
                    status_code=500, detail="Falta 'openpyxl' para leer Excel (.xlsx)."
                )

            from io import BytesIO

            wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
            ws = wb.active
            it = ws.iter_rows(values_only=True)
            header = next(it, None)
            if not header:
                raise HTTPException(status_code=400, detail=_("schemas.empty_excel"))

            hmap = {str(h).strip().lower(): i for i, h in enumerate(header) if h is not None}
            required = {
                "table_name",
                "type",
                "column_name",
                "data_type",
                "is_primary",
                "is_nullable",
            }
            if not required.issubset(hmap.keys()):
                raise HTTPException(
                    status_code=400,
                    detail=_("schemas.invalid_excel_headers"),
                )

            tmp = _dd(list)
            for r in it:
                if not r:
                    continue

                def val(k):
                    i = hmap[k]
                    return "" if i >= len(r) or r[i] is None else str(r[i]).strip()

                tname = val("table_name")
                ttype = val("type").lower()
                cname = val("column_name")
                ctype = val("data_type").upper()
                ip = val("is_primary").lower()
                inull = val("is_nullable").lower()

                if not tname or ttype not in ("origin", "destination"):
                    continue
                if not cname or ctype not in ALLOWED_TYPES:
                    continue

                tmp[(tname, ttype)].append(
                    {
                        "name": cname,
                        "type": ctype,
                        "is_primary": ip in TRUES,
                        "is_nullable": inull in TRUES,
                    }
                )

            for (tname, ttype), cols in tmp.items():
                tables_payload.append({"name": tname, "type": ttype, "columns": cols})

        else:
            raise HTTPException(
                status_code=400,
                detail=_("schemas.unsupported_format"),
            )

        if not tables_payload:
            raise HTTPException(status_code=400, detail=_("schemas.no_valid_tables"))

        # -----------------------------
        #   PRE-FLIGHT LIMITS Free/Pro
        # -----------------------------
        # ✅ Tablas existentes - usar CRUD con desencriptación
        existing_origin = schemas_crud.get_tables_by_project_orm(
            db, project_id, "origin", encryption_key
        )
        existing_dest = schemas_crud.get_tables_by_project_orm(
            db, project_id, "destination", encryption_key
        )
        all_existing_tables = existing_origin + existing_dest
        existing_tables = {
            (t["name"], t["type"]): t["id"]
            for t in all_existing_tables
            if t.get("name")  # Solo incluir tablas con nombre desencriptado
        }

        # ¿cuántas nuevas distintas?
        desired_keys: Set[tuple] = {(str(t["name"]), str(t["type"])) for t in tables_payload}

        # -----------------------------
        #    REPLACE MODE: Eliminar tablas duplicadas
        # -----------------------------
        tables_replaced = 0
        columns_replaced = 0

        if duplicate_strategy == "replace":
            # Identificar tablas a reemplazar (existen Y están en el archivo)
            tables_to_replace = []
            for tname, ttype in desired_keys:
                tid = existing_tables.get((tname, ttype))
                if tid:
                    tables_to_replace.append((tid, tname, ttype))

            # Eliminar tablas existentes (CASCADE eliminará columnas)
            for tid, tname, ttype in tables_to_replace:
                # Obtener conteo de columnas antes de eliminar
                col_count_row = _fetch_one(
                    db, "SELECT COUNT(*) AS count FROM columns WHERE table_id = :tid", {"tid": tid}
                )
                col_count = int(col_count_row["count"]) if col_count_row else 0

                # Eliminar tabla (schemas_crud maneja la eliminación de columnas y mappings en cascada)
                table_info, deleted_cols, mappings_del = schemas_crud.delete_table_orm(
                    db, tid, current_user["id"]
                )
                if table_info:
                    tables_replaced += 1
                    columns_replaced += col_count
                    # Remover de existing_tables para que sea tratada como nueva
                    existing_tables.pop((tname, ttype), None)

            db.commit()  # Commit de eliminaciones antes de proceder con inserciones

        # ¿cuántas nuevas distintas?
        new_needed = sum(1 for k in desired_keys if k not in existing_tables)

        # Cuenta actual
        current_count = int(
            _fetch_one(
                db, "SELECT COUNT(*) AS c FROM tables WHERE project_id = :pid", {"pid": project_id}
            )["c"]
        )

        if current_count + new_needed > max_tables:
            # Límite de número de tablas
            return redirect_with_flash(_("schemas.max_pro") if has_pro else _("schemas.max_free"))

        # ✅ Límite de columnas por tabla - cargar con desencriptación
        exist_cols_map = {}
        if all_existing_tables:
            for table in all_existing_tables:
                if table.get("name") and table.get("id"):
                    # Cargar columnas desencriptadas para cada tabla
                    columns = schemas_crud.get_columns_by_table_orm(db, table["id"], encryption_key)
                    col_names = {col["name"] for col in columns if col.get("name")}
                    exist_cols_map[(table["name"], table["type"])] = col_names

        # Para cada tabla del payload, suma existentes + nuevas no duplicadas
        from collections import defaultdict as dd

        incoming_cols = _dd(set)
        for t in tables_payload:
            key = (t["name"], t["type"])
            for c in t.get("columns") or []:
                cname = c["name"].strip()
                if cname:
                    incoming_cols[key].add(cname)

        for key, new_cols in incoming_cols.items():
            current_cols = exist_cols_map.get(key, set())
            to_add = {c for c in new_cols if c not in current_cols}
            if len(current_cols) + len(to_add) > MAX_COLS:
                return redirect_with_flash(
                    _("attributes.max_col_pro") if has_pro else _("attributes.max_col_free")
                )

        # -----------------------------
        #       INSERCIÓN EN BD
        # -----------------------------
        table_key_to_id = dict(existing_tables)

        # Tablas únicas a insertar
        unique_table_keys = []
        seen = set()
        tables_duplicates_skipped = 0
        for t in tables_payload:
            key = (t["name"], t["type"])
            if key not in seen:
                if key not in table_key_to_id:
                    seen.add(key)
                    unique_table_keys.append((project_id, t["name"], t["type"]))
                else:
                    # Tabla ya existe, contar como duplicado omitido
                    tables_duplicates_skipped += 1
                    seen.add(key)

        total_tables_payload = len(desired_keys)
        total_columns_payload = sum(len(t.get("columns") or []) for t in tables_payload)

        tables_added = 0
        if unique_table_keys:
            #  Usar schemas_crud con soporte de encriptación
            def _insert_tables_with_encryption(batch):
                # batch: [(project_id, name, type), ...]
                created = []
                for pid, name, ttype in batch:
                    try:
                        table = schemas_crud.create_table_orm(
                            db=db,
                            project_id=pid,
                            name=name,
                            type=ttype,
                            user_id=current_user["id"],
                            encryption_key=encryption_key,
                        )
                        created.append(table)
                    except Exception as e:
                        # Handle conflicts (table already exists) - skip and continue
                        import logging

                        logger = logging.getLogger(__name__)
                        logger.debug(f"Table {name} ({ttype}) already exists or error: {e}")
                        continue
                return created

            BATCH = 300
            for i in range(0, len(unique_table_keys), BATCH):
                batch = unique_table_keys[i : i + BATCH]
                rows = await run_in_threadpool(_insert_tables_with_encryption, batch)
                for r in rows:
                    # Use the decrypted name that's already in the result
                    table_key_to_id[(r["name"], r["type"])] = r["id"]
                tables_added += len(rows)
                db.commit()

            # Cargar ids que cayeron en DO NOTHING usando CRUD (con desencriptación)
            missing = [t for t in unique_table_keys if (t[1], t[2]) not in table_key_to_id]
            if missing:
                missing_names_by_type: Dict[str, Set] = {}
                for mt in missing:
                    ttype = str(mt[2])  # 'origin' or 'destination'
                    if ttype not in missing_names_by_type:
                        missing_names_by_type[ttype] = set()
                    missing_names_by_type[ttype].add(mt[1])  # name

                # Get tables by type using CRUD
                for ttype, names in missing_names_by_type.items():
                    fetched_tables = schemas_crud.get_tables_by_project_orm(
                        db, project_id, table_type=ttype, encryption_key=encryption_key
                    )
                    for table in fetched_tables:
                        if table.get("name") in names:
                            table_key_to_id[(table["name"], table["type"])] = table["id"]

        # Construir set de table_ids afectados
        affected_table_ids = set(
            table_key_to_id.get((t["name"], t["type"])) for t in tables_payload
        )
        affected_table_ids.discard(None)

        # Columns existentes por table_id
        from collections import defaultdict

        existing_cols_by_tid = _dd(set)
        if affected_table_ids:
            for r in _fetch_all(
                db,
                "SELECT c.table_id, c.name FROM columns c WHERE c.table_id = ANY(CAST(:ids AS uuid[]))",
                {"ids": list(affected_table_ids)},
            ):
                existing_cols_by_tid[r["table_id"]].add(r["name"])

        # Preparar columnas a insertar
        to_insert_cols = []
        duplicates_skipped = 0
        for t in tables_payload:
            tid = table_key_to_id.get((t["name"], t["type"]))
            if not tid:
                #  Usar CRUD para buscar tabla (con desencriptación)
                fetched_tables = schemas_crud.get_tables_by_project_orm(
                    db, project_id, table_type=str(t["type"]), encryption_key=encryption_key
                )
                table_obj = next(
                    (tbl for tbl in fetched_tables if tbl.get("name") == t["name"]), None
                )
                if table_obj:
                    tid = table_obj["id"]
                    table_key_to_id[(t["name"], t["type"])] = tid
                else:
                    continue

            existing = existing_cols_by_tid.get(tid, set())
            for c in t.get("columns") or []:
                cname = c["name"]
                ctype = c["type"]
                if cname in existing:
                    duplicates_skipped += 1
                    continue
                to_insert_cols.append(
                    (
                        tid,
                        cname,
                        ctype,
                        t["type"],
                        bool(c["is_primary"]),
                        bool(c["is_nullable"]),
                        current_time,
                        current_time,
                    )
                )
                existing.add(cname)

        columns_added = 0
        if to_insert_cols:
            #  Usar schemas_crud con soporte de encriptación
            def _insert_cols_with_encryption(batch):
                # batch: [(table_id, name, data_type, table_type, is_primary, is_nullable, created_at, updated_at), ...]
                created_count = 0
                for tid, cname, ctype, ttype, is_pk, is_null, ct, ut in batch:
                    try:
                        column = schemas_crud.create_column_orm(
                            db=db,
                            table_id=tid,
                            name=cname,
                            data_type=ctype,
                            table_type=ttype,
                            is_nullable=is_null,
                            is_primary=is_pk,
                            is_autogenerated=False,
                            user_id=current_user["id"],
                            encryption_key=encryption_key,
                        )
                        created_count += 1
                    except Exception as e:
                        # Handle conflicts (column already exists) - skip and continue
                        import logging

                        logger = logging.getLogger(__name__)
                        logger.debug(f"Column {cname} already exists or error: {e}")
                        continue
                return created_count

            BATCHC = 2000
            for i in range(0, len(to_insert_cols), BATCHC):
                col_batch = to_insert_cols[i : i + BATCHC]
                count = await run_in_threadpool(_insert_cols_with_encryption, col_batch)
                db.commit()
                columns_added += count

        #  Counters corregidos: reflejan inserciones REALES en DB
        counters = {
            "tables_added": tables_added,  # Inserciones reales en DB
            "columns_added": columns_added,  # Inserciones reales en DB
            "duplicates_skipped": (
                duplicates_skipped if duplicate_strategy == "skip" else 0
            ),  # Columnas duplicadas
            "tables_duplicates_skipped": (
                tables_duplicates_skipped if duplicate_strategy == "skip" else 0
            ),  # Tablas duplicadas
            "tables_replaced": tables_replaced if duplicate_strategy == "replace" else 0,
            "columns_replaced": columns_replaced if duplicate_strategy == "replace" else 0,
            "total_in_file": total_tables_payload,  # Info adicional: cuántas había en archivo
            "total_columns_in_file": total_columns_payload,
        }

        # ========== LOGGING DE SEGURIDAD ==========
        import logging

        security_logger = logging.getLogger("security.file_upload")
        security_logger.info(
            f"FILE_UPLOAD_SUCCESS: user_id={current_user['id']}, "
            f"project_id={project_id}, filename={safe_filename}, "
            f"format={format}, file_size={len(content)}, "
            f"tables_added={tables_added}, columns_added={columns_added}, "
            f"total_in_file={total_tables_payload}, total_columns_in_file={total_columns_payload}"
        )

        #  VALIDACIÓN POST-IMPORT: Verificar que datos fueron guardados
        if total_tables_payload > 0 and tables_added == 0:
            # Only raise error if we didn't skip everything (all duplicates in skip mode or replaced in replace mode)
            if duplicate_strategy == "skip" and tables_duplicates_skipped == 0:
                # No tables added and no duplicates skipped means encryption issue
                raise HTTPException(
                    status_code=500,
                    detail=_(
                        "La importación aparentemente tuvo éxito pero no se crearon tablas. "
                        "Verifica que el proyecto esté desbloqueado si tiene encriptación habilitada."
                    ),
                )
            elif duplicate_strategy == "replace" and tables_replaced == 0:
                # In replace mode, if no tables added and none replaced, something went wrong
                raise HTTPException(
                    status_code=500,
                    detail=_(
                        "La importación aparentemente tuvo éxito pero no se crearon tablas. "
                        "Verifica que el proyecto esté desbloqueado si tiene encriptación habilitada."
                    ),
                )
            # If tables_duplicates_skipped > 0 or tables_replaced > 0, it's expected that tables_added = 0

        if total_columns_payload > 0 and columns_added == 0:
            # Only raise error if we didn't skip everything
            if duplicates_skipped == 0:  # No duplicates means encryption issue
                raise HTTPException(
                    status_code=500,
                    detail=_(
                        "La importación aparentemente tuvo éxito pero no se crearon columnas. "
                        "Verifica que el proyecto esté desbloqueado si tiene encriptación habilitada."
                    ),
                )
            # If duplicates_skipped > 0, it's expected that columns_added = 0

        # Broadcast + respuesta
        await realtime_broadcast(
            project_id,
            {
                "type": "tables_imported",
                "ts": datetime.utcnow().isoformat(),
                "actor_id": str(current_user["id"]),
                "counters": counters,
                "format": format,
            },
        )
        return JSONResponse({"status": "ok", **counters})

    except HTTPException as he:
        db.rollback()
        # Logging de seguridad para errores HTTP (validaciones fallidas)
        import logging

        security_logger = logging.getLogger("security.file_upload")
        security_logger.warning(
            f"FILE_UPLOAD_VALIDATION_ERROR: user_id={current_user.get('id', 'unknown')}, "
            f"project_id={project_id}, filename={getattr(file, 'filename', 'unknown')}, "
            f"format={format}, error={he.detail}, status_code={he.status_code}"
        )
        raise
    except Exception as e:
        db.rollback()
        # Logging de seguridad para errores internos
        import logging
        import traceback

        security_logger = logging.getLogger("security.file_upload")
        security_logger.error(
            f"FILE_UPLOAD_INTERNAL_ERROR: user_id={current_user.get('id', 'unknown')}, "
            f"project_id={project_id}, filename={getattr(file, 'filename', 'unknown')}, "
            f"format={format}, error={str(e)}"
        )
        print("[IMPORT][ERROR]", e)
        print(traceback.format_exc())

        #  Proporcionar mensajes de error más específicos según el tipo de error
        error_str = str(e).lower()
        if "encrypt" in error_str or "decrypt" in error_str or "cipher" in error_str:
            raise HTTPException(
                status_code=500,
                detail=_(
                    "Error de encriptación durante la importación. "
                    "Verifica que el proyecto esté desbloqueado correctamente y que la sesión no haya expirado."
                ),
            )
        elif "unique" in error_str or "duplicate" in error_str:
            raise HTTPException(
                status_code=409,
                detail=_(
                    "Algunos elementos ya existen en el proyecto. Revisa los datos y vuelve a intentar."
                ),
            )
        elif "permission" in error_str or "access" in error_str:
            raise HTTPException(
                status_code=403,
                detail=_("No tienes permisos suficientes para realizar esta operación."),
            )
        else:
            raise HTTPException(status_code=500, detail=_("schemas.error_generic"))


# =============================================================================
# ENDPOINT: Generar Script de Migración SQL
# =============================================================================


def _create_adapter_from_config(config: DBConnectionConfig):
    """
    Crea un adapter de BD a partir de configuración descifrada.

    Args:
        config: Configuración de conexión con credenciales descifradas

    Returns:
        Instancia del adapter apropiado (PostgreSQLAdapter, MySQLAdapter, etc.)

    Raises:
        ValueError: Si el dialecto no está soportado
    """
    adapter_config = config.to_adapter_config()

    if config.dialect == "postgresql":
        return PostgreSQLAdapter(adapter_config)
    elif config.dialect == "mysql":
        return MySQLAdapter(adapter_config)
    elif config.dialect == "sqlserver":
        return SQLServerAdapter(adapter_config)
    elif config.dialect == "oracle":
        return OracleAdapter(adapter_config)
    else:
        raise ValueError(f"Unsupported dialect: {config.dialect}")


@router.post("/projects/{project_id}/generate-migration-script")
async def generate_migration_script(
    project_id: str,
    db_orm=Depends(get_db_session),
    current_user=Depends(get_current_user),
):
    """
    Genera un script SQL de migración basado en los mappings de un proyecto.

    Este endpoint:
    1. Obtiene todos los mappings del proyecto
    2. Identifica qué sesiones de conexión se necesitan
    3. Valida que las sesiones existan y no hayan expirado
    4. Crea adapters para las bases de datos necesarias
    5. Genera el script SQL resolviendo hashes SHA-256 si es necesario

    Args:
        project_id: UUID del proyecto
        db_orm: Sesión de SQLAlchemy
        current_user: Usuario actual autenticado

    Returns:
        JSON con:
        - script: Script SQL generado
        - mappings_count: Número de mappings procesados
        - generated_at: Timestamp de generación

    Raises:
        HTTP 404 Not Found: Si el proyecto no existe o no tiene mappings
        HTTP 404 Not Found: Si un hash no se encuentra en la BD externa
        HTTP 410 Gone: Si una sesión expiró
        HTTP 503 Service Unavailable: Si no se puede conectar a una BD externa
    """
    from app.crud import mappings_crud, master_key_crud

    try:
        # Obtener encryption_key de la sesión del usuario
        encryption_key = master_key_crud.get_encryption_key_from_session_orm(
            db_orm, current_user["id"]
        )

        # 1. Obtener todos los mappings del proyecto
        mappings = mappings_crud.get_mappings_by_project_orm(
            db=db_orm, project_id=project_id, encryption_key=encryption_key
        )

        if not mappings:
            raise HTTPException(
                status_code=404, detail=f"No mappings found for project {project_id}"
            )

        # 2. Identificar sesiones necesarias
        origin_session_ids = set()
        dest_session_ids = set()

        for mapping in mappings:
            if mapping.get("origin_source_type") == "connector":
                session_id = mapping.get("origin_session_id")
                if session_id:
                    origin_session_ids.add(session_id)

            if mapping.get("destination_source_type") == "connector":
                session_id = mapping.get("destination_session_id")
                if session_id:
                    dest_session_ids.add(session_id)

        # 3. Obtener proyecto name (para el header del script)
        project_row = _fetch_one(
            db_orm, "SELECT name FROM projects WHERE id = :pid", {"pid": project_id}
        )
        project_name = project_row["name"] if project_row else None

        # 4. Crear DBConnectionService
        redis_client = get_redis_client()
        master_key_str = os.getenv("ENCRYPTION_MASTER_KEY")
        if not master_key_str:
            raise HTTPException(status_code=500, detail="Encryption master key not configured")
        db_service = DBConnectionService(redis_client, master_key_str)

        # 5. Crear adapters para sesiones necesarias
        origin_adapter = None
        dest_adapter = None

        # Adapter para origen (si hay al menos un mapping connector en origen)
        if origin_session_ids:
            # Usar la primera sesión (asumiendo que todos los mappings usan la misma BD origen)
            origin_session_id = list(origin_session_ids)[0]

            try:
                session = db_service.get_session(current_user["id"], origin_session_id)

                # Validar ownership
                if session.user_id != current_user["id"]:
                    raise HTTPException(
                        status_code=403, detail="Origin session does not belong to current user"
                    )

                # Descifrar configuración
                config = db_service.get_decrypted_config(session)

                # Crear adapter
                origin_adapter = _create_adapter_from_config(config)

                # Conectar
                origin_adapter.connect()

            except SessionNotFoundError:
                raise HTTPException(
                    status_code=404, detail=f"Origin session {origin_session_id} not found"
                )
            except SessionExpiredError:
                raise HTTPException(
                    status_code=410,
                    detail=f"Origin session {origin_session_id} has expired. Please reconnect to the database.",
                )
            except DBConnectionError as e:
                raise HTTPException(
                    status_code=503, detail=f"Cannot connect to origin database: {str(e)}"
                )

        # Adapter para destino (si hay al menos un mapping connector en destino)
        if dest_session_ids:
            # Usar la primera sesión
            dest_session_id = list(dest_session_ids)[0]

            try:
                session = db_service.get_session(current_user["id"], dest_session_id)

                # Validar ownership
                if session.user_id != current_user["id"]:
                    raise HTTPException(
                        status_code=403,
                        detail="Destination session does not belong to current user",
                    )

                # Descifrar configuración
                config = db_service.get_decrypted_config(session)

                # Crear adapter
                dest_adapter = _create_adapter_from_config(config)

                # Conectar
                dest_adapter.connect()

            except SessionNotFoundError:
                raise HTTPException(
                    status_code=404, detail=f"Destination session {dest_session_id} not found"
                )
            except SessionExpiredError:
                raise HTTPException(
                    status_code=410,
                    detail=f"Destination session {dest_session_id} has expired. Please reconnect to the database.",
                )
            except DBConnectionError as e:
                raise HTTPException(
                    status_code=503, detail=f"Cannot connect to destination database: {str(e)}"
                )

        # 6. Generar script SQL
        try:
            script = MigrationScriptService.generate_insert_script(
                mappings=mappings,
                origin_adapter=origin_adapter,
                dest_adapter=dest_adapter,
                project_name=project_name,
            )

        except HashNotFoundError as e:
            raise HTTPException(
                status_code=404,
                detail=f"Hash {e.hash_value[:8]}... not found in {e.side} database. "
                f"Table or column may have been renamed or deleted.",
            )
        except MigrationScriptServiceError as e:
            raise HTTPException(
                status_code=500, detail=f"Error generating migration script: {str(e)}"
            )

        # 7. Retornar script
        return {
            "script": script,
            "mappings_count": len(mappings),
            "generated_at": datetime.utcnow().isoformat(),
        }

    except HTTPException:
        raise

    except Exception as e:
        logger.error(f"Error generating migration script: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Unexpected error: {str(e)}")

    finally:
        # Desconectar adapters
        try:
            if origin_adapter and origin_adapter.is_connected:
                origin_adapter.disconnect()
        except Exception as e:
            logger.warning(f"Error disconnecting origin adapter: {str(e)}")

        try:
            if dest_adapter and dest_adapter.is_connected:
                dest_adapter.disconnect()
        except Exception as e:
            logger.warning(f"Error disconnecting destination adapter: {str(e)}")
